import openpyxl
import json

# Открываем JSON файл
with open('schedule.json', 'r', encoding='utf-8') as json_file:
    json = json.load(json_file)

# Создаём новую книгу Excel и активируем лист
book = openpyxl.Workbook()
sheet = book.active

# Заполняем заголовки
sheet['A1'] = 'ID'
sheet['B1'] = 'DAY'
sheet['C1'] = 'EVERY_WEEK'
sheet['D1'] = 'ODD_WEEK'
sheet['E1'] = 'NUMBER'
sheet['F1'] = 'TIME'
sheet['G1'] = 'ROOM'
sheet['H1'] = 'LESSON'
sheet['I1'] = 'TEACHER'
sheet['J1'] = 'LOCATION'
# Начальная строка для записи данных
row = 2
# Заполняем строки данными(вычленяем данные)
for schedule in json['schedule']:
    day = schedule['day']
    every_week = schedule.get('every_week', False)
    odd_week = schedule.get('odd_week', False)
    # Заполняем каждую строку и столбец
    for subject in schedule['subjects']:
        sheet.cell(row=row, column=1, value=json['id'])  # ID из верхнего уровня
        sheet.cell(row=row, column=2, value=day)
        sheet.cell(row=row, column=3, value=every_week)
        sheet.cell(row=row, column=4, value=odd_week)
        sheet.cell(row=row, column=5, value=subject['number'])
        sheet.cell(row=row, column=6, value=subject['time'])
        sheet.cell(row=row, column=7, value=subject['room'])
        sheet.cell(row=row, column=8, value=subject['lesson'])
        sheet.cell(row=row, column=9, value=subject['teacher'])
        sheet.cell(row=row, column=10, value=subject['location'])
        row += 1
book.save('shedule_xls.xlsx')
